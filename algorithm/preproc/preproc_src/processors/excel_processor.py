import pandas as pd
import openpyxl
from typing import Dict, List, Any, Tuple
import numpy as np
from datetime import datetime, date
import logging
import os

logger = logging.getLogger(__name__)


class ExcelProcessor:
    def __init__(self):
        # 更新支持的文件格式
        self.supported_formats = ['.xlsx', '.xls', '.csv']

    def _convert_cell_value(self, value):
        """转换单元格值，确保适合MySQL插入"""
        if pd.isna(value):  # 处理NaN值
            return ''
        elif isinstance(value, (pd.Timestamp, datetime)):
            # 转换日期时间类型为字符串
            return value.strftime('%Y-%m-%d %H:%M:%S')
        elif isinstance(value, date):
            # 转换日期类型为字符串
            return value.strftime('%Y-%m-%d')
        elif isinstance(value, (int, float, np.integer, np.floating)):
            # 转换数值类型为字符串
            return str(value)
        else:
            # 其他类型直接转换为字符串
            return str(value)

    def process_excel(self, file_path: str) -> Dict[str, Any]:
        """
        处理Excel或CSV文件，返回结构化数据
        返回格式: {
            'file_name': str,
            'sheets': [
                {
                    'sheet_name': str,
                    'data': List[List],  # 原始数据
                    'columns': List[str],  # 列名
                    'metadata': Dict
                }
            ]
        }
        """
        # 获取文件扩展名
        _, file_ext = os.path.splitext(file_path)
        file_ext = file_ext.lower()

        # 根据文件类型选择处理方法
        if file_ext == '.csv':
            return self._process_csv(file_path)
        elif file_ext in ['.xlsx', '.xls']:
            return self._process_excel_file(file_path)
        else:
            logger.error(f"不支持的文件格式: {file_ext}")
            return {'error': f'不支持的文件格式: {file_ext}'}

    def _process_excel_file(self, file_path: str) -> Dict[str, Any]:
        """处理Excel文件 (.xlsx, .xls)"""
        try:
            # 使用pandas读取Excel
            excel_data = pd.ExcelFile(file_path)
            result = {
                'file_name': os.path.basename(file_path),
                'sheets': [],
                'total_sheets': len(excel_data.sheet_names),
                'processed_at': datetime.now().isoformat(),
                'file_type': 'excel'
            }

            for sheet_name in excel_data.sheet_names:
                # 读取每个sheet
                df = pd.read_excel(file_path, sheet_name=sheet_name)

                # 处理当前sheet
                sheet_info = self._process_dataframe(df, sheet_name)
                result['sheets'].append(sheet_info)

            return result

        except Exception as e:
            logger.error(f"使用pandas处理Excel失败: {e}")
            return self._fallback_excel_process(file_path)

    def _process_csv(self, file_path: str) -> Dict[str, Any]:
        """处理CSV文件"""
        try:
            # 尝试自动检测分隔符和编码
            result = {
                'file_name': os.path.basename(file_path),
                'sheets': [],
                'total_sheets': 1,  # CSV只有一个"sheet"
                'processed_at': datetime.now().isoformat(),
                'file_type': 'csv'
            }

            # 首先尝试读取前几行来检测分隔符
            try:
                # 尝试常见分隔符
                for delimiter in [',', ';', '\t', '|']:
                    try:
                        df = pd.read_csv(file_path, delimiter=delimiter, nrows=5, encoding='utf-8')
                        if df.shape[1] > 1:  # 如果成功分割出多列
                            # 重新读取整个文件
                            df = pd.read_csv(file_path, delimiter=delimiter, encoding='utf-8')
                            break
                    except:
                        continue
                else:
                    # 如果常见分隔符都不行，尝试自动检测
                    df = pd.read_csv(file_path, sep=None, engine='python', encoding='utf-8')
            except UnicodeDecodeError:
                # 尝试其他编码
                for encoding in ['latin1', 'iso-8859-1', 'cp1252', 'gbk', 'gb2312']:
                    try:
                        df = pd.read_csv(file_path, encoding=encoding)
                        break
                    except:
                        continue
                else:
                    # 所有编码都失败
                    df = pd.read_csv(file_path, encoding='utf-8', errors='replace')

            # 处理CSV数据（当作单个sheet）
            sheet_name = os.path.basename(file_path).replace('.csv', '')
            sheet_info = self._process_dataframe(df, sheet_name)
            result['sheets'].append(sheet_info)

            return result

        except Exception as e:
            logger.error(f"处理CSV文件失败: {e}")
            return self._fallback_csv_process(file_path)

    def _process_dataframe(self, df: pd.DataFrame, sheet_name: str) -> Dict[str, Any]:
        """通用DataFrame处理逻辑"""
        # 转换列名为字符串
        columns = [str(col) for col in df.columns.tolist()]

        # 处理数据：转换每个单元格值
        data = []
        for _, row in df.iterrows():
            row_data = []
            for cell in row:
                row_data.append(self._convert_cell_value(cell))
            data.append(row_data)

        return {
            'sheet_name': sheet_name,
            'columns': columns,
            'data': data,
            'shape': df.shape,
            'metadata': {
                'row_count': df.shape[0],
                'col_count': df.shape[1],
                'data_types': {col: str(df[col].dtype) for col in df.columns}
            }
        }

    def _fallback_excel_process(self, file_path: str) -> Dict[str, Any]:
        """使用openpyxl作为Excel备选方案"""
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            result = {
                'file_name': os.path.basename(file_path),
                'sheets': [],
                'total_sheets': len(wb.sheetnames),
                'processed_at': datetime.now().isoformat(),
                'file_type': 'excel'
            }

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                data = []

                # 获取最大行和列
                max_row = ws.max_row
                max_col = ws.max_column

                # 读取数据
                for row in ws.iter_rows(min_row=1, max_row=max_row,
                                        min_col=1, max_col=max_col,
                                        values_only=True):
                    row_data = []
                    for cell in row:
                        row_data.append(self._convert_cell_value(cell))
                    data.append(row_data)

                if data:
                    columns = [str(cell) for cell in data[0]] if max_row > 0 else []
                    sheet_data = data[1:] if max_row > 1 else []
                else:
                    columns = []
                    sheet_data = []

                sheet_info = {
                    'sheet_name': sheet_name,
                    'columns': columns,
                    'data': sheet_data,
                    'shape': (len(sheet_data), len(columns)),
                    'metadata': {
                        'row_count': len(sheet_data),
                        'col_count': len(columns)
                    }
                }

                result['sheets'].append(sheet_info)

            return result

        except Exception as e:
            logger.error(f"使用openpyxl处理Excel失败: {e}")
            return {'error': str(e), 'file_name': os.path.basename(file_path)}

    def _fallback_csv_process(self, file_path: str) -> Dict[str, Any]:
        """使用Python内置csv模块作为CSV备选方案"""
        import csv

        try:
            result = {
                'file_name': os.path.basename(file_path),
                'sheets': [],
                'total_sheets': 1,
                'processed_at': datetime.now().isoformat(),
                'file_type': 'csv'
            }

            data = []
            # 尝试不同编码
            encodings = ['utf-8', 'latin1', 'iso-8859-1', 'cp1252', 'gbk']

            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        # 尝试不同分隔符
                        dialect = csv.Sniffer().sniff(f.read(1024))
                        f.seek(0)
                        reader = csv.reader(f, dialect)
                        data = list(reader)
                        break
                except:
                    continue

            if not data:
                # 如果自动检测失败，使用默认逗号分隔符
                with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
                    reader = csv.reader(f)
                    data = list(reader)

            if data:
                columns = [str(cell) for cell in data[0]] if data else []
                sheet_data = data[1:] if len(data) > 1 else []
            else:
                columns = []
                sheet_data = []

            sheet_name = os.path.basename(file_path).replace('.csv', '')
            sheet_info = {
                'sheet_name': sheet_name,
                'columns': columns,
                'data': sheet_data,
                'shape': (len(sheet_data), len(columns)),
                'metadata': {
                    'row_count': len(sheet_data),
                    'col_count': len(columns)
                }
            }

            result['sheets'].append(sheet_info)
            return result

        except Exception as e:
            logger.error(f"使用csv模块处理CSV失败: {e}")
            return {'error': str(e), 'file_name': os.path.basename(file_path)}

    def generate_sql_schema(self, excel_data: Dict[str, Any]) -> List[Dict[str, str]]:
        """
        为Excel数据生成SQL建表语句
        """
        schemas = []

        for sheet in excel_data['sheets']:
            sheet_name = sheet['sheet_name']
            columns = sheet['columns']

            # 清理表名和列名（移除特殊字符）
            clean_table_name = f"excel_{excel_data['file_name'].replace('.', '_')}_{sheet_name}"
            clean_table_name = ''.join(c for c in clean_table_name if c.isalnum() or c == '_')

            # 生成SQL
            sql_columns = []
            for col in columns:
                clean_col = ''.join(c for c in str(col) if c.isalnum() or c == '_')
                sql_columns.append(f"`{clean_col}` TEXT")

            create_table_sql = f"""
            CREATE TABLE IF NOT EXISTS `{clean_table_name}` (
                id INT AUTO_INCREMENT PRIMARY KEY,
                {', '.join(sql_columns)},
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
            """

            schemas.append({
                'table_name': clean_table_name,
                'sql': create_table_sql,
                'sheet_name': sheet_name
            })

        return schemas